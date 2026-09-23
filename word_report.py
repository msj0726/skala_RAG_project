"""Create an editable Word companion from the rendered report Markdown."""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_font(style, size, color=None, bold=False):
    style.font.name = "AppleGothic"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(*color)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "AppleGothic")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def render(markdown, target):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.9)
    doc.core_properties.title = "KV cache 최적화 기술 시장성 및 도메인 평가 보고서"
    for name, size, color, bold in [
        ("Normal", 10, (36, 48, 68), False),
        ("Heading 1", 18, (16, 44, 67), True),
        ("Heading 2", 12, (18, 109, 128), True),
    ]:
        set_font(doc.styles[name], size, color, bold)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.18
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.text = "KV CACHE | MLA & CXL-PNM"
    footer.style = doc.styles["Normal"]
    footer.runs[0].font.size = Pt(8)

    lines = markdown.read_text(encoding="utf-8").splitlines()
    i = 0
    chapter = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("# "):
            if chapter > 1:
                doc.add_page_break()
            chapter += 1
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("| "):
            rows = []
            while True:
                rows.append([c.strip() for c in line.strip().strip("|").split("|")])
                if i >= len(lines) or not lines[i].strip().startswith("| "):
                    break
                line = lines[i].strip()
                i += 1
            rows = [r for r in rows if not all(c == "---" for c in r)]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = ([3.0, 6.5, 7.8] if len(rows[0]) == 3 else [4.0, 13.3])
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.width = Cm(widths[c])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.08
                        for run in p.runs:
                            run.font.name = "AppleGothic"
                            run.font.size = Pt(8.2)
                            run.font.bold = r == 0
                            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "AppleGothic")
                    shade(cell, "E2F1F4" if r == 0 else "F7FAFB" if r % 2 == 0 else "FFFFFF")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line[2:-2])
            run.bold = True
            run.font.color.rgb = RGBColor(18, 109, 128)
        else:
            doc.add_paragraph(line)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    render(args.output_dir / "report.md", args.output_dir / "report.docx")
